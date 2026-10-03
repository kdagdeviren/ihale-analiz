# -*- coding: utf-8 -*-
"""
ECZANE İHALE ANALİZ PROGRAMI
=============================
Tek dosyalik masaustu uygulamasi. Kullanici istedigi sayida yillik hareket
dosyasini ekler, "Tamamlandi" der, program yillari kendisi tanir ve
karsilastirmali Excel raporunu uretir.

BAGIMLILIK: yalnizca openpyxl (+ .xls dosyalari icin xlrd).
pandas ve numpy KULLANILMAZ - boylece Wine icinde sorunsuz derlenir.
"""

import calendar
import datetime as dt
import os
import re
import statistics
import sys
import threading
import traceback

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
    TK_VAR = True
except Exception:
    TK_VAR = False


# =============================================================================
#  AYARLAR
# =============================================================================
FIILEN_STOKSUZ_ESIK = 1.0
FIILEN_STOKSUZ_MIN_GUN = 3
NEGATIF_DONEMDE_CIKIS_VARSA_STOKTA_SAY = True
SIFIR_CIKISLI_ILACLAR_DAHIL = False

PROGRAM_ADI = "DEÜ Hastane Eczanesi Planlama Birimi İhale Analiz Programı"
SURUM = "1.0"
IMZA = "Yusuf Kağan DAĞDEVİREN ® 2026"

AYLAR = {1: "OCAK", 2: "ŞUBAT", 3: "MART", 4: "NİSAN", 5: "MAYIS", 6: "HAZİRAN",
         7: "TEMMUZ", 8: "AĞUSTOS", 9: "EYLÜL", 10: "EKİM", 11: "KASIM", 12: "ARALIK"}
AY_ADI = {1: "Ocak", 2: "Şubat", 3: "Mart", 4: "Nisan", 5: "Mayıs", 6: "Haziran",
          7: "Temmuz", 8: "Ağustos", 9: "Eylül", 10: "Ekim", 11: "Kasım", 12: "Aralık"}


# =============================================================================
#  GENEL YARDIMCILAR
# =============================================================================
def _norm_baslik(c):
    s = str(c).strip()
    for a, b in [("İ", "I"), ("ı", "i"), ("Ş", "S"), ("ş", "s"), ("Ğ", "G"),
                 ("ğ", "g"), ("Ü", "U"), ("ü", "u"), ("Ö", "O"), ("ö", "o"),
                 ("Ç", "C"), ("ç", "c")]:
        s = s.replace(a, b)
    return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_").upper()


def _to_num(x):
    """'2.945' -> 2945.0 ; 2945 -> 2945.0 ; None -> 0.0"""
    if x is None:
        return 0.0
    if isinstance(x, bool):
        return 0.0
    if isinstance(x, (int, float)):
        try:
            if x != x:          # NaN
                return 0.0
        except Exception:
            pass
        return float(x)
    s = str(x).strip().replace("\xa0", "").replace(" ", "")
    if s in ("", "-", "--", "None", "nan"):
        return 0.0
    neg = s.startswith("-") or (s.startswith("(") and s.endswith(")"))
    s = s.strip("()").lstrip("+-")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") \
            else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", "") if re.fullmatch(r"\d{1,3}(,\d{3})+", s) else s.replace(",", ".")
    elif "." in s and re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        v = float(s)
    except ValueError:
        return 0.0
    return -v if neg else v


def _kod_temizle(x, uyari=None):
    if x is None:
        return ""
    if isinstance(x, int):
        return str(x)
    if isinstance(x, float):
        if x != x:
            return ""
        if x.is_integer():
            return str(int(x))
        s = f"{x:.10f}".rstrip("0")
        if re.fullmatch(r"\d{1,3}\.\d{3}", s):
            if uyari is not None:
                uyari.append("Bazı etken madde kodları ondalıklı okundu; binlik ayracı "
                             "kabul edilip birleştirildi.")
            return s.replace(".", "")
        return s
    s = str(x).strip()
    if re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    elif re.fullmatch(r"\d{1,3}(,\d{3})+", s):
        s = s.replace(",", "")
    return s


TR_ALFABE = " 0123456789ABCÇDEFGĞHIİJKLMNOÖPQRSŞTUÜVWXYZ"
_TR_SIRA = {ch: i for i, ch in enumerate(TR_ALFABE)}


def _tr_upper(s):
    return str(s).replace("i", "İ").replace("ı", "I").upper()


def tr_key(s):
    return [_TR_SIRA.get(ch, 500 + ord(ch)) for ch in _tr_upper(s)]


def _vir(x, b=1):
    return f"{x:.{b}f}".replace(".", ",")


def _sutun_bul(kolonlar, adaylar, zorunlu=True, icerik=None):
    for a in adaylar:
        if a in kolonlar:
            return kolonlar[a]
    if icerik:
        for norm, idx in kolonlar.items():
            if icerik in norm:
                return idx
    if zorunlu:
        raise KeyError("Sütun bulunamadı: " + " / ".join(adaylar))
    return None


# =============================================================================
#  TARIH COZUMLEME
# =============================================================================
_TARIH_KALIP = re.compile(r"(\d{1,4})[.\-/](\d{1,2})[.\-/](\d{2,4})")


def _tarih_parcala(deger):
    """Doner: (a, b, yil) ham parcalar; gun/ay sirasi henuz belirsiz."""
    if isinstance(deger, dt.datetime):
        return (deger.day, deger.month, deger.year, True)
    if isinstance(deger, dt.date):
        return (deger.day, deger.month, deger.year, True)
    s = str(deger).strip()
    m = _TARIH_KALIP.search(s)
    if not m:
        return None
    a, b, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if len(m.group(1)) == 4:              # YYYY-MM-DD
        return (y, b, a, True) if False else (int(m.group(3)), b, a, True)
    if y < 100:
        y += 2000 if y < 70 else 1900
    return (a, b, y, False)


# =============================================================================
#  DOSYA OKUMA (openpyxl / xlrd)
# =============================================================================
def _satirlar_xlsx(yol, log):
    wb = load_workbook(yol, read_only=True, data_only=True)
    ilk_bas = None
    for sayfa in wb.sheetnames:
        ws = wb[sayfa]
        bas = None
        sayac = 0
        for satir in ws.iter_rows(values_only=True):
            if bas is None:
                if satir is None or all(h is None for h in satir):
                    continue
                bas = tuple(_norm_baslik(h) if h is not None else "" for h in satir)
                if ilk_bas is None:
                    ilk_bas = bas
                elif bas != ilk_bas:
                    break
                continue
            sayac += 1
            yield bas, satir
        if bas is not None and sayac and len(wb.sheetnames) > 1:
            log(f"      sayfa '{sayfa}': {sayac:,} satır")
    wb.close()


def _satirlar_xls(yol, log):
    import xlrd
    kitap = xlrd.open_workbook(yol)
    ilk_bas = None
    for ws in kitap.sheets():
        if ws.nrows < 2:
            continue
        bas = tuple(_norm_baslik(ws.cell_value(0, c)) for c in range(ws.ncols))
        if ilk_bas is None:
            ilk_bas = bas
        elif bas != ilk_bas:
            continue
        for r in range(1, ws.nrows):
            satir = []
            for c in range(ws.ncols):
                v = ws.cell_value(r, c)
                if ws.cell_type(r, c) == xlrd.XL_CELL_DATE:
                    try:
                        y, mo, d, *_ = xlrd.xldate_as_tuple(v, kitap.datemode)
                        v = dt.datetime(y, mo, d)
                    except Exception:
                        pass
                satir.append(v)
            yield bas, tuple(satir)
        if len(kitap.sheets()) > 1:
            log(f"      sayfa '{ws.name}': {ws.nrows - 1:,} satır")


def satirlari_oku(yol, log):
    if yol.lower().endswith(".xls"):
        return _satirlar_xls(yol, log)
    return _satirlar_xlsx(yol, log)


# =============================================================================
#  VERIYI HAZIRLA
# =============================================================================
def veri_hazirla(yol, log, uyari):
    """Doner: kayitlar listesi [(kod, ad, tarih, giris, cikis, kson, devir)]"""
    ham = []
    kolonlar = None
    ix = {}
    gunfirst_oy = [0, 0]      # [ay-once dogru, gun-once dogru]

    for bas, satir in satirlari_oku(yol, log):
        if kolonlar != bas:
            kolonlar = bas
            harita = {}
            for i, h in enumerate(bas):
                harita.setdefault(h, i)
            ix = {
                "tarih": _sutun_bul(harita, ["DAY_ID", "TARIH", "GUN"], icerik="DAY"),
                "kod": _sutun_bul(harita, ["ETKEN_MADDE_KODU", "ETKEN_MADDE_KOD"], icerik="KOD"),
                "ad": _sutun_bul(harita, ["ETKEN_MADDE", "ETKEN_MADDE_ADI"]),
                "gun": _sutun_bul(harita, ["AYIN_GUNU", "AYIN_GUN"], False, "AYIN"),
                "devir": _sutun_bul(harita, ["DEVIR"], False),
                "giris": _sutun_bul(harita, ["GIRIS"]),
                "cikis": _sutun_bul(harita, ["CIKIS"]),
                "kson": _sutun_bul(harita, ["KALAN_SON"]),
            }

        n = len(satir)
        if ix["kod"] >= n or ix["tarih"] >= n:
            continue
        kod = _kod_temizle(satir[ix["kod"]], uyari)
        ad = satir[ix["ad"]] if ix["ad"] < n else None
        ad = str(ad).strip() if ad is not None else ""
        if not kod or kod in ("nan", "None") or not ad or ad.lower() in ("nan", "none"):
            continue

        p = _tarih_parcala(satir[ix["tarih"]])
        if p is None:
            continue
        a, b, yil, kesin = p
        ayin_gunu = None
        if ix["gun"] is not None and ix["gun"] < n:
            g = _to_num(satir[ix["gun"]])
            ayin_gunu = int(g) if g > 0 else None
        if not kesin and ayin_gunu:
            if a == ayin_gunu:
                gunfirst_oy[1] += 1
            if b == ayin_gunu:
                gunfirst_oy[0] += 1

        ham.append((kod, ad, (a, b, yil, kesin), ayin_gunu,
                    _to_num(satir[ix["giris"]]) if ix["giris"] < n else 0.0,
                    _to_num(satir[ix["cikis"]]) if ix["cikis"] < n else 0.0,
                    _to_num(satir[ix["kson"]]) if ix["kson"] < n else 0.0,
                    _to_num(satir[ix["devir"]]) if (ix["devir"] is not None
                                                    and ix["devir"] < n) else 0.0))

    gun_once = gunfirst_oy[1] > gunfirst_oy[0]

    kayitlar = []
    bozuk = 0
    for kod, ad, (a, b, yil, kesin), ayin_gunu, gi, ci, ks, dv in ham:
        if kesin:
            gun, ay = a, b
        else:
            gun, ay = (a, b) if gun_once else (b, a)
            if ayin_gunu and gun != ayin_gunu and 1 <= ayin_gunu <= 31:
                gun, ay = ayin_gunu, (b if gun_once else a)
        try:
            tarih = dt.date(yil, ay, gun)
        except ValueError:
            bozuk += 1
            continue
        kayitlar.append((kod, ad, tarih, gi, ci, ks, dv))

    if bozuk:
        uyari.append(f"{bozuk:,} satırın tarihi çözümlenemedi; atlandı.")
    atlanan = len(ham) - len(kayitlar)
    return kayitlar


def kapsam_belirle(kayitlar, uyari, log):
    yil_say = {}
    for k in kayitlar:
        yil_say[k[2].year] = yil_say.get(k[2].year, 0) + 1
    yil = max(yil_say, key=yil_say.get)
    if len(yil_say) > 1:
        uyari.append("Dosyada birden fazla yıl var; en çok satıra sahip yıl "
                     f"({yil}) hesaplandı.")
    kayitlar = [k for k in kayitlar if k[2].year == yil]

    gun_say = {}
    for k in kayitlar:
        gun_say[k[2]] = gun_say.get(k[2], 0) + 1
    gunler = sorted(gun_say)
    if len(gunler) > 10:
        medyan = statistics.median(gun_say.values())
        while len(gunler) > 1 and gun_say[gunler[-1]] < 0.5 * medyan:
            kesilen = gunler[-1]
            uyari.append(f"{kesilen.strftime('%d.%m.%Y')} günü yarım çekilmiş "
                         f"({gun_say[kesilen]} satır, normal ~{int(medyan)}); "
                         f"hesap dışı bırakıldı.")
            kayitlar = [k for k in kayitlar if k[2] < kesilen]
            gunler.pop()

    while True:
        aylar_var = {}
        for k in kayitlar:
            aylar_var.setdefault(k[2].month, set()).add(k[2])
        if len(aylar_var) <= 1:
            break
        son_ay = max(aylar_var)
        if len(aylar_var[son_ay]) >= 15:
            break
        uyari.append(f"{AY_ADI[son_ay]} ayında yalnızca {len(aylar_var[son_ay])} günlük "
                     f"veri var; yarım ay tabloyu yanıltmasın diye hesap dışı bırakıldı.")
        kayitlar = [k for k in kayitlar if k[2].month < son_ay]

    son_ay = max(k[2].month for k in kayitlar)
    son_tarih = max(k[2] for k in kayitlar)
    ilk_tarih = min(k[2] for k in kayitlar)
    log(f"      dönem: {ilk_tarih.strftime('%d.%m.%Y')} – "
        f"{son_tarih.strftime('%d.%m.%Y')}  ({len(kayitlar):,} satır)")
    return kayitlar, yil, son_ay, son_tarih


def gunluk_ozet(kayitlar, uyari):
    birlesik = {}
    isim_say = {}
    mukerrer = False
    for kod, ad, tarih, gi, ci, ks, dv in kayitlar:
        isim_say.setdefault(kod, {})
        isim_say[kod][ad] = isim_say[kod].get(ad, 0) + 1
        anahtar = (kod, tarih)
        if anahtar in birlesik:
            mukerrer = True
            e = birlesik[anahtar]
            birlesik[anahtar] = (e[0] + gi, e[1] + ci, ks, e[3])
        else:
            birlesik[anahtar] = (gi, ci, ks, dv)
    if mukerrer:
        uyari.append("Aynı ilaç-aynı gün için birden fazla kayıt var; "
                     "GİRİŞ/ÇIKIŞ toplandı, KALAN_SON'da son satır alındı.")

    ilac = {}
    for (kod, tarih), (gi, ci, ks, dv) in birlesik.items():
        ilac.setdefault(kod, {})[tarih] = (gi, ci, ks, dv)
    isim = {k: max(v, key=v.get) for k, v in isim_say.items()}
    return ilac, isim


# =============================================================================
#  HESAPLAMA
# =============================================================================
def yil_hesapla(ilac, isim, yil, son_ay, son_tarih):
    aylar = list(range(1, son_ay + 1))
    ay_gun = {m: calendar.monthrange(yil, m)[1] for m in aylar}
    satirlar = []

    for kod, kayit in ilac.items():
        gunluk, stok_onceki = [], None
        for m in aylar:
            bitti = False
            for d in range(1, ay_gun[m] + 1):
                t = dt.date(yil, m, d)
                if t > son_tarih:
                    bitti = True
                    break
                if t in kayit:
                    gi, ci, ks, dv = kayit[t]
                    gb = stok_onceki if stok_onceki is not None else dv
                else:
                    gi = ci = 0.0
                    gb = stok_onceki
                    ks = stok_onceki
                gunluk.append([m, gb, gi, ci, ks, t in kayit])
                stok_onceki = ks
            if bitti:
                break

        vardi = [False] * len(gunluk)
        for i, (m, gb, gi, ci, ks, kay) in enumerate(gunluk):
            if kay:
                v = (gb is not None and gb > 0) or (ks is not None and ks > 0) or ci > 0
                if gi > 0:
                    v = True
            else:
                v = gb is not None and gb > 0
            vardi[i] = v

        if NEGATIF_DONEMDE_CIKIS_VARSA_STOKTA_SAY:
            i = 0
            while i < len(gunluk):
                if gunluk[i][4] is not None and gunluk[i][4] < 0:
                    j = i
                    while j + 1 < len(gunluk) and gunluk[j + 1][4] is not None \
                            and gunluk[j + 1][4] < 0:
                        j += 1
                    ci_idx = [k for k in range(i, j + 1) if gunluk[k][3] > 0]
                    if ci_idx:
                        for k in range(min(ci_idx), max(ci_idx) + 1):
                            vardi[k] = True
                    i = j + 1
                else:
                    i += 1

        poz = [x[3] for x in gunluk if x[3] > 0]
        tipik = float(statistics.median(poz)) if poz else 0.0
        esik = FIILEN_STOKSUZ_ESIK * tipik
        fiilen = 0
        if esik > 0:
            i = 0
            while i < len(gunluk):
                m, gb, gi, ci, ks, kay = gunluk[i]
                if vardi[i] and gi <= 0 and ci <= 0 and gb is not None and 0 < gb < esik:
                    j = i
                    while j + 1 < len(gunluk):
                        m2, gb2, gi2, ci2, ks2, k2 = gunluk[j + 1]
                        if vardi[j + 1] and gi2 <= 0 and ci2 <= 0 and gb2 is not None \
                                and 0 < gb2 < esik:
                            j += 1
                        else:
                            break
                    if (j - i + 1) >= FIILEN_STOKSUZ_MIN_GUN:
                        for k in range(i, j + 1):
                            vardi[k] = False
                            fiilen += 1
                    i = j + 1
                else:
                    i += 1

        adet = {m: 0.0 for m in aylar}
        gun = {m: 0 for m in aylar}
        for i, (m, gb, gi, ci, ks, kay) in enumerate(gunluk):
            adet[m] += ci
            if vardi[i]:
                gun[m] += 1
        toplam = sum(adet.values())
        tg = sum(gun.values())
        satirlar.append({"KOD": kod, "AD": isim.get(kod, ""), "adet": adet, "gun": gun,
                         "toplam": toplam, "toplam_gun": tg, "fiilen": fiilen,
                         "tipik": tipik,
                         "ortalama": (toplam / tg * 30) if tg else 0.0})

    satirlar.sort(key=lambda r: (tr_key(r["AD"]), r["KOD"]))
    return {"yil": yil, "son_ay": son_ay, "son_tarih": son_tarih, "aylar": aylar,
            "ay_gun": ay_gun, "satirlar": satirlar,
            "donem_gun": sum(min(ay_gun[m], son_tarih.day if m == son_tarih.month
                                 else ay_gun[m]) for m in aylar)}


def yil_isle(yol, log, uyari):
    log(f"   • {os.path.basename(yol)}")
    kayitlar = veri_hazirla(yol, log, uyari)
    log(f"      okundu: {len(kayitlar):,} geçerli satır")
    kayitlar, yil, son_ay, son_tarih = kapsam_belirle(kayitlar, uyari, log)
    ilac, isim = gunluk_ozet(kayitlar, uyari)
    sonuc = yil_hesapla(ilac, isim, yil, son_ay, son_tarih)
    hareketli = sum(1 for r in sonuc["satirlar"] if r["toplam"] > 0)
    log(f"      {yil} yılı: {len(sonuc['satirlar']):,} kalem, {hareketli:,} tanesi hareketli")
    return sonuc


# =============================================================================
#  MOLEKUL / DOZ AYRISTIRICI
# =============================================================================
def mnorm(s):
    s = str(s).upper()
    for a, b in [("İ", "I"), ("Ş", "S"), ("Ğ", "G"), ("Ü", "U"), ("Ö", "O"),
                 ("Ç", "C"), ("Â", "A"), ("Î", "I")]:
        s = s.replace(a, b)
    return " ".join(s.split())


TUZ = {"SODYUM", "HIDROKLORUR", "HCL", "BESILAT", "MALEAT", "SULFAT", "TROMETAMOL",
       "DISOPROKSIL", "SUKSINAT", "ASETAT", "FOSFAT", "TARTARAT", "MESILAT",
       "SITRAT", "LAKTAT", "DIHIDRAT", "MONOHIDRAT", "POTASYUM", "KALSIYUM"}
FORM = {"FLAKON", "AMPUL", "TABLET", "KAPSUL", "SURUP", "NEBUL", "SOLUSYON", "COZELTI",
        "SUSPANSIYON", "DAMLA", "KREM", "POMAT", "POMAD", "MERHEM", "JEL", "SPREY",
        "TORBA", "SISE", "LAVMAN", "SASE", "OVUL", "FITIL", "GARGARA", "INHALER",
        "KALEM", "ENJEKTOR", "TOZ", "GRANUL", "EFERVESAN", "DRAJE", "PASTIL", "SET",
        "OFTALMIK", "IM", "IV", "SC", "MR", "NEBULIZATOR", "INFUZYONLUK", "ENJEKSIYON"}
TOPIKAL = ("KREM", "POMAT", "POMAD", "MERHEM", "JEL", "LOSYON", "SAMPUAN")
BIRIM_CARPAN = {"MG": 1.0, "GR": 1000.0, "G": 1000.0}

_SAYI = r"\d{1,3}(?:\.\d{3})+|\d+(?:[.,]\d+)?"
_BIRIMLER = r"MCG|MG|GR|G|UNITE|IU|ML|CC"
KONS = re.compile(rf"({_SAYI})\s*({_BIRIMLER})\s*/\s*({_SAYI})?\s*(?:ML|CC)\b")
DUZ = re.compile(rf"({_SAYI})\s*({_BIRIMLER})\b")
HACIM = re.compile(rf"({_SAYI})\s*(?:ML|CC)\b")
ORAN = re.compile(rf"({_SAYI})\s*/\s*({_SAYI})\s*({_BIRIMLER})\b")
BIRIM_SINIF = {"MG": ("MG", 1.0), "GR": ("MG", 1000.0), "G": ("MG", 1000.0),
               "MCG": ("MCG", 1.0), "IU": ("IU", 1.0), "UNITE": ("IU", 1.0),
               "ML": ("ML", 1.0), "CC": ("ML", 1.0)}


def _say(t):
    t = t.strip()
    if re.fullmatch(r"\d{1,3}(\.\d{3})+", t):
        return float(t.replace(".", ""))
    return float(t.replace(",", "."))


def doz_coz(ad):
    u = mnorm(ad)
    m = KONS.search(u)
    if m and m.group(2) not in ("ML", "CC"):
        sinif, carp = BIRIM_SINIF[m.group(2)]
        a = _say(m.group(1)) * carp
        b = _say(m.group(3)) if m.group(3) else 1.0
        h = HACIM.search(u[m.end():])
        if h:
            c = _say(h.group(1))
            if abs(c - b) > 1e-9:
                return (a * c / b, sinif)
        return (a, sinif)
    m = ORAN.search(u)
    if m and m.group(3) not in ("ML", "CC"):
        sinif, carp = BIRIM_SINIF[m.group(3)]
        a, b = _say(m.group(1)), _say(m.group(2))
        v = (a if a > b else a / b) * carp
        if v > 0:
            return (v, sinif)
    for m in DUZ.finditer(u):
        sinif, carp = BIRIM_SINIF[m.group(2)]
        v = _say(m.group(1)) * carp
        if v > 0:
            return (v, sinif)
    return None


def topikal_mi(ad):
    return any(t in mnorm(ad).split() for t in TOPIKAL)


def molekul(ad):
    u = re.sub(r"^(KT|ASI)\s+", "", mnorm(ad).strip())
    cikan = []
    for p in u.split():
        t = p.strip(".,()")
        if not t:
            continue
        if re.match(r"^[\d%]", t) or re.match(r"^[.,/%]+$", t):
            break
        if t in FORM or t in ("MG", "GR", "G", "ML", "MCG", "UNITE", "IU", "CC"):
            break
        if t == "+":
            cikan.append("+")
            continue
        if t in TUZ and cikan:
            continue
        cikan.append(t)
    return " ".join(cikan).strip()


TARIF_ONEK = ("BES.", "SER.", "SERR", "NONIYONIK", "MAKROSIKLIK", "GADO")


def grup_disi(ad):
    return mnorm(ad).startswith(TARIF_ONEK) or len(molekul(ad)) < 4



# =============================================================================
#  ORTALAMA GUVENILIRLIK NOTU
# =============================================================================
def ortalama_notu(r, aylar, tam_gun):
    tg, toplam, ort = r["toplam_gun"], r["toplam"], r["ortalama"]
    if tg <= 0 or toplam <= 0:
        return 0, ""
    stoklu = [m for m in aylar if r["gun"][m] > 0]
    hiz = {m: r["adet"][m] / r["gun"][m] * 30 for m in stoklu}
    seviye, notlar = 0, []
    if tg < 30:
        seviye = 3
        notlar.append(f"Yalnızca {tg} gün stokta kaldı; 30 güne çarpım ortalamayı "
                      f"yaklaşık {30/tg:.0f} kat büyütüyor. Ham çıkışa ({toplam:,.0f}) bakın.")
    elif tg < 0.40 * tam_gun:
        seviye = max(seviye, 1)
        notlar.append(f"Dönemin {tam_gun} gününün yalnızca {tg} günü stokta.")
    en_buyuk = max(aylar, key=lambda m: r["adet"][m])
    pay = r["adet"][en_buyuk] / toplam
    if pay >= 0.70 and len(stoklu) >= 2:
        seviye = max(seviye, 2)
        notlar.append(f"Toplam çıkışın %{pay*100:.0f} kadarı tek ayda ({AY_ADI[en_buyuk]}).")
    if len(hiz) >= 2 and ort > 0:
        ey = max(hiz, key=hiz.get)
        if hiz[ey] >= 3 * ort:
            seviye = max(seviye, 2)
            notlar.append(f"En yoğun ay {AY_ADI[ey]}, ortalamanın {_vir(hiz[ey]/ort)} katı.")
    bos = [m for m in stoklu if r["adet"][m] == 0]
    if len(bos) >= 3:
        seviye = max(seviye, 1)
        notlar.append(f"{len(bos)} ayda stokta olduğu halde hiç çıkış olmamış.")
    if r["fiilen"] >= 15:
        seviye = max(seviye, 2)
        notlar.append(f"{r['fiilen']} gün donmuş bakiye (tipik günlük çıkış "
                      f"{r['tipik']:.0f}); bu günler stoksuz sayıldı.")
    return seviye, " ".join(notlar)


# =============================================================================
#  EXCEL - TEK YIL SAYFASI
# =============================================================================
KALIN_BEYAZ = Font(bold=True, color="FFFFFF")
INCE = Side(style="thin", color="BFBFBF")
SEV_DOLGU = {1: PatternFill("solid", fgColor="FFF2CC"),
             2: PatternFill("solid", fgColor="FCE4D6"),
             3: PatternFill("solid", fgColor="F8CBAD")}


def yil_sayfasi(wb, sonuc, ad):
    ws = wb.create_sheet(ad)
    aylar = sonuc["aylar"]
    bas = ["Etken Madde Kodu", "Etken Madde Adı"]
    for m in aylar:
        bas += [AYLAR[m], "Çıkan Gün"]
    bas += ["Toplam Çıkış", "Toplam Gün", "Aylık Ortalama", "Fiilen Stoksuz Gün",
            "Ortalama Notu"]
    ws.append(bas)

    tam_gun = sonuc["donem_gun"]
    veriler = [r for r in sonuc["satirlar"]
               if SIFIR_CIKISLI_ILACLAR_DAHIL or r["toplam"] > 0]
    for r in veriler:
        r["seviye"], r["not"] = ortalama_notu(r, aylar, tam_gun)
        satir = [r["KOD"], r["AD"]]
        for m in aylar:
            satir += [int(round(r["adet"][m])), r["gun"][m]]
        satir += [int(round(r["toplam"])), r["toplam_gun"],
                  int(round(r["ortalama"])), r["fiilen"], r["not"]]
        ws.append(satir)

    for j, b in enumerate(bas, 1):
        h = ws.cell(row=1, column=j)
        h.font = KALIN_BEYAZ
        h.fill = PatternFill("solid", fgColor="404040" if j <= 2 else
                             ("2E75B6" if b == "Çıkan Gün" else "1F4E78"))
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(bas))}{ws.max_row}"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 52
    for j in range(3, len(bas)):
        ws.column_dimensions[get_column_letter(j)].width = 14
    ws.column_dimensions[get_column_letter(len(bas))].width = 78
    for i, r in enumerate(veriler, start=2):
        for j in range(3, len(bas)):
            c = ws.cell(row=i, column=j)
            c.number_format = "#,##0"
            c.alignment = Alignment(horizontal="center")
        ws.cell(row=i, column=len(bas)).alignment = Alignment(wrap_text=True, vertical="top")
        if r["seviye"]:
            ws.cell(row=i, column=len(bas)).fill = SEV_DOLGU[r["seviye"]]
            ws.cell(row=i, column=len(bas) - 2).fill = SEV_DOLGU[r["seviye"]]
    return ws


# =============================================================================
#  BIRLESIK TABLO (N YIL)
# =============================================================================
def birlesik_kur(yillar):
    """yillar: yil_hesapla ciktilarinin listesi. Yeniden eskiye siralanir."""
    yillar = sorted(yillar, key=lambda s: -s["yil"])
    etiket = [str(s["yil"]) for s in yillar]
    ix = {}
    for s in yillar:
        ix[str(s["yil"])] = {r["KOD"]: r for r in s["satirlar"]}

    kodlar = set()
    for s in yillar:
        kodlar |= {r["KOD"] for r in s["satirlar"]}

    isim = {}
    for s in sorted(yillar, key=lambda x: x["yil"]):
        for r in s["satirlar"]:
            isim[r["KOD"]] = r["AD"]

    kayit = []
    for k in kodlar:
        d = {"kod": k, "ad": isim[k]}
        for e in etiket:
            r = ix[e].get(k)
            d[f"c{e}"] = r["toplam"] if r else 0.0
            d[f"g{e}"] = r["toplam_gun"] if r else 0
            d[f"f{e}"] = r["fiilen"] if r else 0
            d[f"o{e}"] = r["ortalama"] if r else 0.0
        d["cT"] = sum(d[f"c{e}"] for e in etiket)
        d["gT"] = sum(d[f"g{e}"] for e in etiket)
        d["fT"] = sum(d[f"f{e}"] for e in etiket)
        d["oT"] = (d["cT"] / d["gT"] * 30) if d["gT"] else 0.0
        if d["cT"] <= 0:
            continue
        # ardisik degisimler (eskiden yeniye)
        eski_yeni = list(reversed(etiket))
        d["degisim"] = []
        for a, b in zip(eski_yeni, eski_yeni[1:]):
            oa, ob = d[f"o{a}"], d[f"o{b}"]
            d["degisim"].append((f"{a}→{b}", (ob / oa - 1) * 100 if oa > 0 else None))
        if len(eski_yeni) > 2:
            oa, ob = d[f"o{eski_yeni[0]}"], d[f"o{eski_yeni[-1]}"]
            d["degisim"].append((f"{eski_yeni[0]}→{eski_yeni[-1]}",
                                 (ob / oa - 1) * 100 if oa > 0 else None))
        kayit.append(d)

    return kayit, etiket, yillar


def durum_ve_not(d, etiket, donem_gun):
    yeni, eski = etiket[0], etiket[-1]
    var = {e: d[f"c{e}"] > 0 for e in etiket}
    notlar, sev = [], 0

    if not var[yeni] and any(var[e] for e in etiket[1:]):
        sev = 3
        kaynak = next(e for e in etiket[1:] if var[e])
        s = f"{yeni} yılında hiç çıkış yok"
        s += (f", ancak {int(d['g'+yeni])} gün stokta durmuş (talep yok ya da kullanılmamış)"
              if d[f"g{yeni}"] > 0 else " ve hiç stoğa girmemiş (temin edilememiş olabilir)")
        notlar.append(f"{s}. {kaynak} yılında aylık {d['o'+kaynak]:,.0f} seviyesindeydi. "
                      f"İhale listesinden düşmeden önce klinikle teyit edin.")
        durum = f"{yeni}'da yok"
    elif var[yeni] and not any(var[e] for e in etiket[1:]):
        sev = max(sev, 2)
        notlar.append(f"Sadece {yeni} yılında hareket görmüş; yeni kalem. Çok yıllık "
                      f"ortalama yerine {yeni} verisi ve klinik öngörü esas alınmalı.")
        durum = f"Yeni ({yeni})"
    elif var[yeni] and not var[eski]:
        sev = max(sev, 1)
        notlar.append(f"{eski} yılında hiç hareket yok; ortalama daha kısa bir döneme dayanıyor.")
        durum = "Sonradan girmiş"
    elif var[yeni] and var[eski] and any(not var[e] for e in etiket[1:-1]):
        sev = max(sev, 2)
        kesik = [e for e in etiket[1:-1] if not var[e]]
        notlar.append(f"{', '.join(kesik)} yılında hiç çıkış yok, öncesinde ve sonrasında var. "
                      f"Temin kesintisi olabilir; nedenini araştırın.")
        durum = "Arada kesinti"
    else:
        durum = "Tüm yıllarda var"

    dl = [v for _, v in d["degisim"][:len(etiket) - 1] if v is not None]
    if len(dl) >= 2:
        if all(x >= 25 for x in dl):
            sev = max(sev, 2)
            notlar.append("Üst üste artış trendi; ihale miktarını düz ortalamayla değil, "
                          "trendi ağırlıklandırarak belirleyin.")
        elif all(x <= -25 for x in dl):
            sev = max(sev, 2)
            notlar.append("Üst üste düşüş trendi; formülerden çıkıyor ya da yerini başka "
                          "bir kaleme bırakıyor olabilir.")
    if dl and abs(dl[-1]) >= 60:
        sev = max(sev, 1)
        notlar.append(f"Son dönemde aylık ortalama %{dl[-1]:+.0f} değişmiş; nedenini kontrol edin.")

    dusuk = [e for e in etiket if d[f"c{e}"] > 0 and d[f"g{e}"] < 0.60 * donem_gun.get(e, 1)]
    if len(dusuk) >= 2:
        sev = max(sev, 2)
        notlar.append(f"{', '.join(sorted(dusuk))} yıllarında dönemin %60'ından az süre "
                      f"stokta; kronik stoksuzluk, gerçek talep bu rakamların üzerinde.")
    if d["fT"] >= 15:
        sev = max(sev, 2)
        dok = ", ".join(f"{e}: {int(d['f'+e])}" for e in etiket if d[f"f{e}"] > 0)
        notlar.append(f"Toplam {int(d['fT'])} gün donmuş bakiye ({dok}); kayıtta 'var' "
                      f"görünse de servise dağıtılamamış demektir.")
    if 0 < d["gT"] < 60:
        sev = 3
        notlar.append(f"Toplamda yalnızca {int(d['gT'])} gün stokta; ortalama güvenilir değil, "
                      f"ham çıkışa ({d['cT']:,.0f}) bakın.")
    for et, v in d["degisim"]:
        if v is not None and abs(v) >= 300:
            baz = d[f"o{et.split('→')[0]}"]
            if 0 < baz < 5:
                sev = max(sev, 1)
                notlar.append(f"{et} yüzdesi çok küçük bir tabandan hesaplandı "
                              f"(aylık ort. {baz:.1f}); yüzdeye değil ham rakamlara bakın.")
                break
    return durum, sev, " ".join(notlar)


def molekul_analizi(kayit, etiket):
    yeni, eski = etiket[0], etiket[-1]
    for d in kayit:
        ad = d["ad"]
        d["mol"] = "" if grup_disi(ad) else molekul(ad)
        doz = doz_coz(ad)
        d["doz"], d["birim"] = doz if doz else (None, None)
        d["kt"] = mnorm(ad).startswith("KT ")
        for e in etiket + ["T"]:
            o = d[f"o{e}"] if e != "T" else d["oT"]
            d[f"em{e}"] = o if d["kt"] else (o * d["doz"] if d["doz"] else None)

    gruplar = {}
    for d in kayit:
        if not d["mol"] or d["emT"] is None:
            continue
        birim = (d["birim"] or "MG") + ("-TOPİKAL" if topikal_mi(d["ad"]) else "")
        gruplar.setdefault((d["mol"], birim), []).append(d)

    ozet = []
    for (mol, birim), uyeler in gruplar.items():
        if len(uyeler) < 2:
            continue
        g = {"mol": mol, "birim": birim, "n": len(uyeler),
             "kalemler": sorted(u["ad"] for u in uyeler)}
        for e in etiket + ["T"]:
            g[f"em{e}"] = sum(u[f"em{e}"] for u in uyeler)
        gd = (g[f"em{yeni}"] / g[f"em{eski}"] - 1) * 100 if g[f"em{eski}"] > 0 else None
        g["gd"] = gd
        hareket = sorted(((u[f"em{yeni}"] - u[f"em{eski}"], u) for u in uyeler),
                         key=lambda x: x[0])
        g["dusen"] = [h[1] for h in hareket if h[0] < 0]
        g["yukselen"] = [h[1] for h in reversed(hareket) if h[0] > 0]
        ozet.append(g)

        for u in uyeler:
            u["grup"] = g
            ud = (u[f"em{yeni}"] / u[f"em{eski}"] - 1) * 100 if u[f"em{eski}"] > 0 else None
            pay = (u[f"em{yeni}"] / g[f"em{yeni}"] * 100) if g[f"em{yeni}"] else 0
            u["mol_pay"] = (f"%{pay:.1f}" if 0 < pay < 1 else f"%{pay:.0f}") \
                if g[f"em{yeni}"] else ""
            if ud is None or gd is None:
                continue
            m = None
            if ud <= -35 and abs(gd) <= 20 and g["yukselen"]:
                m = (f"SUNUM KAYMASI: {mol} molekülünün {g['n']} sunumu var; molekül "
                     f"bazında toplam etkin madde %{gd:+.0f} (fiilen sabit). Bu kalemdeki "
                     f"%{ud:+.0f} düşüş gerçek talep düşüşü değil, kullanım "
                     f"\"{g['yukselen'][0]['ad']}\" sunumuna kaymış.")
            elif ud >= 50 and abs(gd) <= 20 and g["dusen"]:
                m = (f"SUNUM KAYMASI: {mol} molekülünün toplamı %{gd:+.0f} ile sabit; "
                     f"bu kalemdeki %{ud:+.0f} artış \"{g['dusen'][0]['ad']}\" "
                     f"sunumundan devralınan kullanımdan geliyor.")
            elif ud <= -35 and gd <= -35:
                m = (f"{mol} molekülünün tamamı geriliyor (%{gd:+.0f}); bu bir sunum "
                     f"kayması değil, gerçek talep düşüşü.")
            elif ud >= 50 and gd >= 50:
                m = (f"{mol} molekülünün tamamı büyüyor (%{gd:+.0f}); artış gerçek "
                     f"talep artışından geliyor.")
            elif ud <= -35 and (gd - ud) >= 25:
                k = g["yukselen"][0] if g["yukselen"] else None
                m = (f"KISMİ SUNUM KAYMASI: bu kalem %{ud:+.0f} düşerken {mol} molekülünün "
                     f"toplamı yalnızca %{gd:+.0f} gerilemiş." +
                     (f" Kullanımın bir kısmı \"{k['ad']}\" sunumuna geçmiş." if k else ""))
            elif ud >= 50 and (ud - gd) >= 40:
                k = g["dusen"][0] if g["dusen"] else None
                m = (f"KISMİ SUNUM KAYMASI: bu kalem %{ud:+.0f} artarken {mol} molekülünün "
                     f"toplamı %{gd:+.0f} değişmiş." +
                     (f" Büyümenin bir kısmı \"{k['ad']}\" sunumundan devralınmış." if k else ""))
            if m:
                u["mol_not"] = m
                u["seviye"] = max(u.get("seviye", 0), 2)

    for d in kayit:
        d.setdefault("mol_not", "")
        d.setdefault("mol_pay", "")
        d.setdefault("grup", None)
    return ozet


def birlesik_sayfa(wb, kayit, etiket, yillar):
    ws = wb.create_sheet("BIRLESIK", 0)
    son = yillar[0]
    ust = ["", ""]
    alt = ["Etken Madde Kodu", "Etken Madde Adı"]
    for s in yillar:
        e = str(s["yil"])
        bitis = s["son_tarih"]
        tam = (bitis.month == 12 and bitis.day >= 30)
        ust += [f"{e}" if tam else f"{e} ({AY_ADI[1][:3]}-{AY_ADI[bitis.month][:3]})"] * 3
        alt += ["Toplam Çıkış", "Toplam Gün", "Aylık Ortalama"]
    ust += ["BİRLEŞİK"] * 5
    alt += ["Toplam Çıkış", "Toplam Gün", "Fiilen Stoksuz Gün", "Aylık Ortalama",
            "Aylık Ort. (Flakon)"]
    dsay = len(kayit[0]["degisim"]) if kayit else 0
    ust += ["DEĞİŞİM (Aylık Ortalama)"] * dsay
    alt += [f"{et} %" for et, _ in (kayit[0]["degisim"] if kayit else [])]
    ust += ["", ""]
    alt += ["Durum", "Notlar"]
    ust += ["MOLEKÜL (SUNUM KAYMASI)"] * 3
    alt += ["Molekül", "Molekül İçi Pay", "Molekül Notu"]
    ust += [f"{son['yil']} AYLIK DETAY"] * (2 * len(son["aylar"]))
    for m in son["aylar"]:
        alt += [AYLAR[m], "Çıkan Gün"]
    ws.append(ust)
    ws.append(alt)

    son_ix = {r["KOD"]: r for r in son["satirlar"]}
    for d in kayit:
        s = [d["kod"], d["ad"]]
        for e in etiket:
            s += [int(round(d[f"c{e}"])), int(d[f"g{e}"]), int(round(d[f"o{e}"]))]
        s += [int(round(d["cT"])), int(d["gT"]), int(d["fT"]), int(round(d["oT"])),
              d.get("flakon") if d.get("flakon") is not None else ""]
        for _, v in d["degisim"]:
            s.append(round(v, 1) if v is not None else "")
        s += [d["durum"], d["not"], d["mol"] if d["grup"] else "", d["mol_pay"], d["mol_not"]]
        r = son_ix.get(d["kod"])
        for m in son["aylar"]:
            s += ([int(round(r["adet"][m])), r["gun"][m]] if r else [0, 0])
        ws.append(s)

    n = len(alt)
    renk = {"BİRLEŞİK": "7F6000", "DEĞİŞİM (Aylık Ortalama)": "375623",
            "MOLEKÜL (SUNUM KAYMASI)": "5B2C6F"}
    mavi = ["1F4E78", "2E75B6", "8EA9DB", "B4C7E7", "D9E2F3"]
    mi = 0
    j = 1
    while j <= n:
        v = ust[j - 1]
        if v:
            k = j
            while k < n and ust[k] == v:
                k += 1
            ws.merge_cells(start_row=1, start_column=j, end_row=1, end_column=k)
            c = ws.cell(row=1, column=j)
            c.value = v
            c.font = KALIN_BEYAZ
            if v in renk:
                fg = renk[v]
            elif v.endswith("AYLIK DETAY"):
                fg = "404040"
            else:
                fg = mavi[min(mi, len(mavi) - 1)]
                mi += 1
            c.fill = PatternFill("solid", fgColor=fg)
            c.alignment = Alignment(horizontal="center", vertical="center")
            j = k + 1
        else:
            j += 1

    for j, b in enumerate(alt, 1):
        h = ws.cell(row=2, column=j)
        h.font = Font(bold=True)
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.border = Border(bottom=INCE)

    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{get_column_letter(n)}{ws.max_row}"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 52
    for j in range(3, n + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14
    ds = alt.index("Durum") + 1
    ns = alt.index("Notlar") + 1
    ms = alt.index("Molekül") + 1
    mps = alt.index("Molekül İçi Pay") + 1
    mns = alt.index("Molekül Notu") + 1
    ws.column_dimensions[get_column_letter(ds)].width = 18
    ws.column_dimensions[get_column_letter(ns)].width = 82
    ws.column_dimensions[get_column_letter(ms)].width = 24
    ws.column_dimensions[get_column_letter(mps)].width = 16
    ws.column_dimensions[get_column_letter(mns)].width = 84
    deg_s = [alt.index(f"{et} %") + 1 for et, _ in (kayit[0]["degisim"] if kayit else [])]

    mor = PatternFill("solid", fgColor="E4D7F5")
    for i, d in enumerate(kayit, start=3):
        for c in range(3, n + 1):
            if c in (ds, ns, ms, mps, mns):
                continue
            h = ws.cell(row=i, column=c)
            h.number_format = "+#,##0.0;-#,##0.0" if c in deg_s else "#,##0"
            h.alignment = Alignment(horizontal="center")
        for c in (ns, mns):
            ws.cell(row=i, column=c).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=ms).alignment = Alignment(horizontal="left")
        if d["seviye"]:
            for c in (ds, ns):
                ws.cell(row=i, column=c).fill = SEV_DOLGU[d["seviye"]]
        if d["mol_not"]:
            for c in (ms, mps, mns):
                ws.cell(row=i, column=c).fill = mor
    return ws


def molekul_sayfasi(wb, ozet, etiket):
    ws = wb.create_sheet("MOLEKUL_OZET")
    bas = ["Molekül", "Birim", "Sunum Sayısı"] + \
          [f"{e} Aylık (etkin madde)" for e in reversed(etiket)] + \
          [f"{etiket[-1]}→{etiket[0]} %", "Değerlendirme", "Sunumlar"]
    ws.append(bas)
    ozet.sort(key=lambda g: -(g["emT"] or 0))
    for g in ozet:
        gd = g["gd"]
        if gd is None:
            y = "Başlangıç yılında hareket yok; karşılaştırma yapılamıyor."
        elif abs(gd) <= 15:
            y = ("Molekül toplamı sabit. Kalemler arasındaki değişimler sunum kaymasıdır; "
                 "ihalede kalemleri tek tek değil, molekül olarak planlayın.")
        elif gd > 15:
            y = f"Molekül toplamı gerçekten büyümüş (%{gd:+.0f}); ihale miktarını artırın."
        else:
            y = f"Molekül toplamı gerçekten gerilemiş (%{gd:+.0f}); talep düşüşü."
        ws.append([g["mol"], g["birim"], g["n"]] +
                  [round(g[f"em{e}"]) for e in reversed(etiket)] +
                  [round(gd, 1) if gd is not None else "", y, " | ".join(g["kalemler"])])
    for j, w in enumerate([26, 12, 12] + [18] * len(etiket) + [14, 80, 90], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for h in ws[1]:
        h.font = KALIN_BEYAZ
        h.fill = PatternFill("solid", fgColor="5B2C6F")
        h.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(bas))}{ws.max_row}"
    for i in range(2, ws.max_row + 1):
        for j in range(3, 3 + len(etiket) + 2):
            ws.cell(row=i, column=j).number_format = "#,##0"
            ws.cell(row=i, column=j).alignment = Alignment(horizontal="center")
        for j in (len(bas) - 1, len(bas)):
            ws.cell(row=i, column=j).alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def bilgi_sayfasi(wb, yillar, etiket, uyari, kayit):
    ws = wb.create_sheet("BILGI")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 108
    sat = [
        ("Program", f"{PROGRAM_ADI} · ver {SURUM}"),
        ("Hazırlayan", IMZA),
        ("Rapor tarihi", dt.datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Analiz edilen yıllar", ", ".join(reversed(etiket))),
    ]
    for s in yillar:
        sat.append((f"{s['yil']} kapsamı",
                    f"01.01.{s['yil']} – {s['son_tarih'].strftime('%d.%m.%Y')} "
                    f"({s['donem_gun']} gün), {len(s['satirlar']):,} kalem"))
    sat += [
        ("", ""),
        ("Aylık Ortalama", "(Toplam Çıkış / fiilen stokta bulunulan gün) x 30. Yıllar arası "
                           "karşılaştırılabilir TEK metrik budur; hem dönem uzunluğuna hem "
                           "bulunurluk süresine göre normalize edilmiştir."),
        ("Toplam Çıkış", "Yıllar arasında DOĞRUDAN karşılaştırmayın: dönem uzunlukları farklı."),
        ("Toplam Gün", "İlacın fiilen stokta bulunduğu gün sayısı (kullanıldığı gün değil)."),
        ("Fiilen Stoksuz Gün", f"Kalan bakiyenin tipik günlük çıkışın "
                               f"%{FIILEN_STOKSUZ_ESIK*100:.0f}'inden az olduğu ve hiç hareket "
                               f"görmeden en az {FIILEN_STOKSUZ_MIN_GUN} gün üst üste öyle "
                               f"kaldığı günler. Rezerve kalmış, dağıtılamayan bakiye kabul "
                               f"edilir ve Toplam Gün'e DAHİL EDİLMEZ."),
        ("Negatif stok kuralı", "KALAN_SON'un negatif kaldığı dönemde ilaç çıkışı varsa "
                                "(HBYS kayıt gecikmesi), ilk çıkıştan son çıkışa kadarki "
                                "günler stokta sayıldı."),
        ("Molekül analizi", "Aynı etken maddenin farklı dozları ve formları tek grupta "
                            "toplanır. KT kalemlerinde çıkış zaten mg cinsindendir; "
                            "diğerlerinde adet x kalem dozu ile çarpılır. Farklı birim "
                            "sınıfları (MG/MCG/IU/ML) birbirine toplanmaz."),
        ("Renkler", "Kırmızımsı: acil inceleme. Turuncu: trend, kesinti ya da kronik "
                    "stoksuzluk. Sarı: sınırlı veri tabanı. Mor: sunum kayması."),
        ("Toplam kalem", f"{len(kayit):,}"),
        ("", ""),
    ]
    for k, v in sat:
        ws.append([k, v])
    if uyari:
        ws.append(["UYARILAR", ""])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="C00000")
        for u in dict.fromkeys(uyari):
            ws.append(["", u])
    for row in ws.iter_rows(min_col=1, max_col=2):
        if row[0].value:
            row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    return ws


# =============================================================================
#  ANA AKIS
# =============================================================================
def analiz_calistir(dosyalar, cikti_klasor, log):
    uyari = []
    log("=" * 64)
    log(PROGRAM_ADI)
    log(IMZA)
    log("=" * 64)
    log(f"ANALİZ BAŞLADI — {len(dosyalar)} dosya")
    log("=" * 64)

    yillar = []
    for yol in dosyalar:
        try:
            yillar.append(yil_isle(yol, log, uyari))
        except Exception as e:
            log(f"      HATA: {e}")
            raise

    if not yillar:
        raise ValueError("Hiçbir dosya işlenemedi.")

    gorulen = [s["yil"] for s in yillar]
    if len(set(gorulen)) != len(gorulen):
        raise ValueError(f"Aynı yıla ait birden fazla dosya var: {gorulen}. "
                         f"Her yıl için tek dosya ekleyin.")

    log("")
    log(f"Tanınan yıllar: {', '.join(str(y) for y in sorted(gorulen, reverse=True))}")

    wb = Workbook()
    wb.remove(wb.active)

    if len(yillar) > 1:
        log("Yıllar birleştiriliyor…")
        kayit, etiket, yillar_s = birlesik_kur(yillar)
        donem_gun = {str(s["yil"]): s["donem_gun"] for s in yillar_s}
        for d in kayit:
            d["durum"], d["seviye"], d["not"] = durum_ve_not(d, etiket, donem_gun)
        log("Molekül / sunum kayması analizi…")
        ozet = molekul_analizi(kayit, etiket)
        for d in kayit:
            doz = doz_coz(d["ad"]) if mnorm(d["ad"]).startswith("KT ") else None
            d["flakon"] = int(round(d["oT"] / doz[0])) if (doz and d["oT"] > 0) else None
        kayit.sort(key=lambda d: (tr_key(d["ad"]), d["kod"]))
        birlesik_sayfa(wb, kayit, etiket, yillar_s)
        log(f"   birleşik tablo: {len(kayit):,} kalem")
        if ozet:
            molekul_sayfasi(wb, ozet, etiket)
            log(f"   molekül özeti: {len(ozet)} grup")
    else:
        kayit, etiket, yillar_s = [], [str(yillar[0]["yil"])], yillar
        log("Tek yıl yüklendi; karşılaştırma tablosu üretilmedi.")

    for s in sorted(yillar, key=lambda x: -x["yil"]):
        yil_sayfasi(wb, s, f"{s['yil']}_AYLIK")
        log(f"   {s['yil']} aylık sayfası eklendi")

    bilgi_sayfasi(wb, sorted(yillar, key=lambda x: -x["yil"]), etiket, uyari, kayit)

    ad = "ihale_analizi_" + "_".join(str(y) for y in sorted(gorulen)) + \
         dt.datetime.now().strftime("_%Y%m%d_%H%M") + ".xlsx"
    yol = os.path.join(cikti_klasor, ad)
    wb.save(yol)

    log("")
    log("=" * 64)
    log(f"TAMAMLANDI → {yol}")
    log("=" * 64)
    if uyari:
        log("")
        log("UYARILAR:")
        for u in dict.fromkeys(uyari):
            log(f"  • {u}")
    return yol


# =============================================================================
#  HIZLI YIL TANIMA (dosya eklenince listede gostermek icin)
# =============================================================================
def yil_tani(yol, ornek=400):
    """Dosyanin ilk birkac yuz satirindan yilini okur."""
    sayac = {}
    n = 0
    for bas, satir in satirlari_oku(yol, lambda *a: None):
        harita = {}
        for i, h in enumerate(bas):
            harita.setdefault(h, i)
        ti = _sutun_bul(harita, ["DAY_ID", "TARIH", "GUN"], icerik="DAY")
        if ti is None or ti >= len(satir):
            continue
        p = _tarih_parcala(satir[ti])
        if p:
            sayac[p[2]] = sayac.get(p[2], 0) + 1
        n += 1
        if n >= ornek:
            break
    if not sayac:
        raise ValueError("Tarih okunamadı")
    return max(sayac, key=sayac.get)


# =============================================================================
#  ARAYUZ
# =============================================================================
if TK_VAR:

    class Uygulama(tk.Tk):
        def __init__(self):
            super().__init__()
            self.title(f"{PROGRAM_ADI}  ·  ver {SURUM}")
            self.geometry("980x680")
            self.minsize(820, 560)
            self.dosyalar = []
            self.calisiyor = False
            self._arayuz()

        # ---------------------------------------------------------------- UI
        def _arayuz(self):
            ust = tk.Frame(self, bg="#1f3a5f", height=72)
            ust.pack(fill="x")
            ust.pack_propagate(False)
            tk.Label(ust, text="DEÜ Hastane Eczanesi Planlama Birimi",
                     bg="#1f3a5f", fg="#b9cbe4",
                     font=("Segoe UI", 9)).place(x=20, y=10)
            tk.Label(ust, text="İhale Analiz Programı", bg="#1f3a5f", fg="white",
                     font=("Segoe UI", 16, "bold")).place(x=18, y=26)
            tk.Label(ust, text="Yıllık hareket dosyalarını ekleyin, Tamamlandı'ya basın",
                     bg="#1f3a5f", fg="#8fa9c8",
                     font=("Segoe UI", 8)).place(x=20, y=52)

            orta = tk.Frame(self, padx=14, pady=12)
            orta.pack(fill="both", expand=True)

            sol = tk.LabelFrame(orta, text=" Yüklenen dosyalar ", padx=8, pady=8)
            sol.pack(side="left", fill="both", expand=True)

            kol = ("dosya", "yil", "satir")
            self.tablo = ttk.Treeview(sol, columns=kol, show="headings", height=9)
            self.tablo.heading("dosya", text="Dosya")
            self.tablo.heading("yil", text="Yıl")
            self.tablo.heading("satir", text="Durum")
            self.tablo.column("dosya", width=330, anchor="w")
            self.tablo.column("yil", width=70, anchor="center")
            self.tablo.column("satir", width=130, anchor="center")
            self.tablo.pack(fill="both", expand=True)

            dugme = tk.Frame(sol)
            dugme.pack(fill="x", pady=(8, 0))
            self.b_ekle = tk.Button(dugme, text="  + Dosya Ekle  ", command=self.dosya_ekle,
                                    bg="#2e75b6", fg="white", font=("Segoe UI", 10, "bold"),
                                    relief="flat", padx=10, pady=6, cursor="hand2")
            self.b_ekle.pack(side="left")
            self.b_sil = tk.Button(dugme, text="  Seçiliyi Çıkar  ", command=self.dosya_sil,
                                   bg="#d9d9d9", relief="flat", padx=10, pady=6,
                                   cursor="hand2")
            self.b_sil.pack(side="left", padx=6)
            self.b_basla = tk.Button(dugme, text="  ✓ Tamamlandı — Analizi Başlat  ",
                                     command=self.basla, bg="#2e7d4f", fg="white",
                                     font=("Segoe UI", 10, "bold"), relief="flat",
                                     padx=10, pady=6, cursor="hand2", state="disabled")
            self.b_basla.pack(side="right")

            sag = tk.LabelFrame(orta, text=" Çıktı klasörü ", padx=8, pady=8)
            sag.pack(side="left", fill="y", padx=(12, 0))
            self.klasor = tk.StringVar(value=os.path.expanduser("~"))
            tk.Label(sag, textvariable=self.klasor, wraplength=190, justify="left",
                     fg="#444").pack(anchor="w")
            tk.Button(sag, text="Değiştir", command=self.klasor_sec, relief="flat",
                      bg="#d9d9d9", padx=8, pady=4, cursor="hand2").pack(anchor="w", pady=6)

            alt = tk.LabelFrame(self, text=" İşlem günlüğü ", padx=8, pady=6)
            alt.pack(fill="both", expand=True, padx=14, pady=(0, 10))
            self.gunluk = tk.Text(alt, height=14, bg="#111418", fg="#d8e2ec",
                                  font=("Consolas", 9), relief="flat", wrap="word")
            kaydir = tk.Scrollbar(alt, command=self.gunluk.yview)
            self.gunluk.configure(yscrollcommand=kaydir.set)
            kaydir.pack(side="right", fill="y")
            self.gunluk.pack(fill="both", expand=True)

            taban = tk.Frame(self, bg="#eceff3")
            taban.pack(fill="x", side="bottom")
            self.durum = tk.Label(taban, text="Hazır — dosya ekleyerek başlayın",
                                  anchor="w", bg="#eceff3", padx=14, pady=5)
            self.durum.pack(side="left")
            tk.Label(taban, text=f"{IMZA}  ·  ver {SURUM}", anchor="e", bg="#eceff3",
                     fg="#3a3a3a", font=("Segoe UI", 9, "bold"),
                     padx=14, pady=5).pack(side="right")

            self.yaz(PROGRAM_ADI)
            self.yaz(IMZA + "  ·  ver " + SURUM)
            self.yaz("")
            self.yaz("1) '+ Dosya Ekle' ile yıllık hareket dosyalarını seçin.")
            self.yaz("   Kaç yıl eklerseniz analiz o kadar yılı kapsar.")
            self.yaz("   Dosyaların yılı içeriğinden otomatik tanınır.")
            self.yaz("2) Ekleme bitince 'Tamamlandı — Analizi Başlat' düğmesine basın.")
            self.yaz("")

        # ------------------------------------------------------------- utils
        def yaz(self, s=""):
            self.gunluk.insert("end", s + "\n")
            self.gunluk.see("end")
            self.update_idletasks()

        def klasor_sec(self):
            k = filedialog.askdirectory(title="Çıktı klasörü seçin")
            if k:
                self.klasor.set(k)

        def dosya_ekle(self):
            yollar = filedialog.askopenfilenames(
                title="Yıllık hareket dosyalarını seçin",
                filetypes=[("Excel dosyaları", "*.xlsx *.xls"), ("Tümü", "*.*")])
            for y in yollar:
                if y in self.dosyalar:
                    continue
                self.dosyalar.append(y)
                self.tablo.insert("", "end", values=(os.path.basename(y), "…", "okunuyor"))
            if yollar:
                threading.Thread(target=self._yillari_tani, daemon=True).start()
            self._dugme_durum()

        def _yillari_tani(self):
            for i, yol in enumerate(self.dosyalar):
                cocuk = self.tablo.get_children()[i]
                if self.tablo.item(cocuk)["values"][1] not in ("…", ""):
                    continue
                try:
                    yil = yil_tani(yol)
                    self.tablo.item(cocuk, values=(os.path.basename(yol), yil, "hazır"))
                except Exception:
                    self.tablo.item(cocuk, values=(os.path.basename(yol), "?", "okunamadı"))
            self._dugme_durum()

        def dosya_sil(self):
            sec = self.tablo.selection()
            for s in sec:
                i = self.tablo.index(s)
                self.tablo.delete(s)
                del self.dosyalar[i]
            self._dugme_durum()

        def _dugme_durum(self):
            self.b_basla.config(state="normal" if (self.dosyalar and not self.calisiyor)
                                else "disabled")
            self.durum.config(text=f"{len(self.dosyalar)} dosya eklendi — "
                                   f"ekleme bittiyse Tamamlandı'ya basın"
                              if self.dosyalar else "Hazır — dosya ekleyerek başlayın")

        # ------------------------------------------------------------ islem
        def basla(self):
            if not self.dosyalar:
                return
            self.calisiyor = True
            for b in (self.b_ekle, self.b_sil, self.b_basla):
                b.config(state="disabled")
            self.durum.config(text="Analiz çalışıyor, lütfen bekleyin…")
            self.gunluk.delete("1.0", "end")
            threading.Thread(target=self._calistir, daemon=True).start()

        def _calistir(self):
            try:
                yol = analiz_calistir(list(self.dosyalar), self.klasor.get(), self.yaz)
                self.durum.config(text=f"Tamamlandı → {os.path.basename(yol)}")
                messagebox.showinfo("Tamamlandı",
                                    f"Rapor oluşturuldu:\n\n{yol}")
                try:
                    os.startfile(os.path.dirname(yol))
                except Exception:
                    pass
            except Exception as e:
                self.yaz("")
                self.yaz("HATA: " + str(e))
                self.yaz(traceback.format_exc())
                self.durum.config(text="Hata oluştu — günlüğe bakın")
                messagebox.showerror("Hata", str(e))
            finally:
                self.calisiyor = False
                self.b_ekle.config(state="normal")
                self.b_sil.config(state="normal")
                self._dugme_durum()


def main():
    if not TK_VAR:
        print("tkinter bulunamadı. Konsol modunda çalıştırmak için:")
        print("  python app.py dosya1.xlsx dosya2.xlsx ...")
        if len(sys.argv) > 1:
            analiz_calistir(sys.argv[1:], os.getcwd(), print)
        return
    Uygulama().mainloop()


if __name__ == "__main__":
    main()
